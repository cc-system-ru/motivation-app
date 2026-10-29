"""Печать расчётных листков: PDF (половина A4 горизонтально) и XLSX."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _register():
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", FONT))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", FONT_B))
        return "DejaVu", "DejaVu-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def _money(v):
    try:
        return f"{float(v):,.0f}".replace(",", " ") + " ₽"
    except (TypeError, ValueError):
        return "—"


def _coef(v):
    try:
        return f"{float(v):.2f}".replace(".", ",")
    except (TypeError, ValueError):
        return "—"


def build_pdf(buf, employees, period_label):
    font, font_b = _register()
    half_h = landscape(A4)[1] / 2  # половина A4 по горизонтали
    page = (landscape(A4)[0], half_h)

    style_title = ParagraphStyle("t", fontName=font_b, fontSize=10, leading=12)
    style_lbl = ParagraphStyle("l", fontName=font, fontSize=6.5, leading=8)
    style_val = ParagraphStyle("v", fontName=font_b, fontSize=8, leading=9)

    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=8 * mm, rightMargin=8 * mm,
                            topMargin=5 * mm, bottomMargin=4 * mm)
    story = []
    for i, emp in enumerate(employees):
        calc = emp.get("calc") or {}
        comps = calc.get("components", [])
        coefs = calc.get("coefficients", {})

        head = Table([[Paragraph(f"РАСЧЁТНЫЙ ЛИСТОК — {period_label}", style_title),
                       Paragraph(emp["fio"], style_title)]],
                     colWidths=[120 * mm, 140 * mm])
        head.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ]))

        meta = Table([[Paragraph(f"Должность: <b>{emp['position']}</b>", style_lbl),
                       Paragraph(f"Начислено: <b>{_money(emp['total'])}</b>", style_lbl),
                       Paragraph(f"Выплачено: <b>{_money(emp.get('paid'))}</b>", style_lbl),
                       Paragraph(f"К получению: <b>{_money((emp['total'] or 0) - (emp.get('paid') or 0))}</b>", style_lbl)]],
                     colWidths=[70 * mm, 60 * mm, 60 * mm, 70 * mm])
        meta.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ]))

        left = [[Paragraph("<b>Составляющая</b>", style_lbl),
                 Paragraph("<b>Сумма</b>", style_lbl), Paragraph("<b>Комментарий</b>", style_lbl)]]
        for c in comps:
            left.append([Paragraph(c["name"], style_lbl),
                         Paragraph(_money(c["value"]), style_val),
                         Paragraph(str(c.get("note") or ""), style_lbl)])
        left.append([Paragraph("<b>ИТОГО НАЧИСЛЕНО</b>", style_lbl),
                     Paragraph(f"<b>{_money(calc.get('total') or emp['total'])}</b>", style_val),
                     Paragraph("", style_lbl)])

        t_left = Table(left, colWidths=[52 * mm, 26 * mm, 42 * mm])
        t_left.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D0D0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFF2CC")),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))

        right = [[Paragraph("<b>Коэффициент</b>", style_lbl), Paragraph("<b>Значение</b>", style_lbl)]]
        for k, v in coefs.items():
            if k.startswith("Выполнение") or k.startswith("Откл"):
                val = f"{float(v)*100:.1f}%".replace(".", ",")
            else:
                val = _coef(v)
            right.append([Paragraph(k, style_lbl), Paragraph(val, style_val)])
        if len(right) == 1:
            right.append([Paragraph("—", style_lbl), Paragraph("—", style_lbl)])
        t_right = Table(right, colWidths=[52 * mm, 20 * mm])
        t_right.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D0D0")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))

        body = Table([[t_left, t_right]], colWidths=[122 * mm, 74 * mm])
        body.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0)]))

        story += [head, meta, Spacer(1, 1.5 * mm), body]
        if i < len(employees) - 1:
            from reportlab.platypus import PageBreak
            story.append(PageBreak())

    doc.build(story)
    return buf


def build_xlsx(buf, employees, period_label):
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчётные листки"
    bold = Font(bold=True, size=10)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="BFBFBF")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([f"Расчётные листки — {period_label}"])
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.append([])
    header = ["ФИО", "Должность", "Составляющая", "Сумма, ₽", "Комментарий",
              "Коэффициенты", "Значение", "Итого начислено", "Выплачено", "К получению"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(ws.max_row, c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.border = bd
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for emp in employees:
        calc = emp.get("calc") or {}
        comps = calc.get("components", [])
        coefs = list((calc.get("coefficients") or {}).items())
        rows = max(len(comps), len(coefs), 1)
        start = ws.max_row + 1
        for i in range(rows):
            r = start + i
            if i == 0:
                ws.cell(r, 1, emp["fio"])
                ws.cell(r, 2, emp["position"])
                ws.cell(r, 8, emp["total"])
                ws.cell(r, 9, emp.get("paid") or 0)
                ws.cell(r, 10, (emp["total"] or 0) - (emp.get("paid") or 0))
            if i < len(comps):
                ws.cell(r, 3, comps[i]["name"])
                ws.cell(r, 4, comps[i]["value"])
                ws.cell(r, 5, comps[i].get("note") or "")
            if i < len(coefs):
                ws.cell(r, 6, coefs[i][0])
                ws.cell(r, 7, coefs[i][1])
        for r in range(start, start + rows):
            for c in range(1, 11):
                cell = ws.cell(r, c)
                cell.border = bd
                cell.font = Font(size=9)
                if c in (4, 8, 9, 10):
                    cell.number_format = '#,##0 "₽"'
                elif c == 7:
                    try:
                        cell.number_format = "0.00" if abs(float(cell.value)) < 100 else "#,##0.0000"
                    except (TypeError, ValueError):
                        pass
        ws.merge_cells(start_row=start, start_column=1, end_row=start + rows - 1, end_column=1)
        ws.merge_cells(start_row=start, start_column=2, end_row=start + rows - 1, end_column=2)
        for c in (8, 9, 10):
            ws.merge_cells(start_row=start, start_column=c, end_row=start + rows - 1, end_column=c)

    widths = [22, 24, 34, 14, 30, 22, 14, 16, 14, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    wb.save(buf)
    return buf
